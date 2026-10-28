"""엑셀 양식 생성 / 명단 읽기 / 배정 결과 내보내기.

파일 대화상자는 호출부(Api)가 처리하고 여기서는 경로만 받는다.
webview 에 의존하지 않으므로 단독으로 실행·검증할 수 있다.
"""
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from labels import grade_of, parse_label, period_num, role_key

# 과목 저장 값 -> 화면·엑셀 이름 ('자습' 은 '수업'으로 보인다. tab-setup.js 의 subjectLabel 과 같다)
SUBJECT_LABEL = {'자습': '수업'}

def write_template(path, teachers=None):

    wb = Workbook()
    ws = wb.active
    ws.title = "감독명단양식"

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    center_align = Alignment(horizontal='center', vertical='center')
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    ws.append(["순", "이름", "담당교과", "담임학급", "참고사항"])

    # 지금 화면의 감독 명단을 그대로 담는다. 명단이 비어 있으면 입력 예시 한 줄을 둔다.
    if teachers:
        for index, t in enumerate(teachers, start=1):
            ws.append([index, t.get('name', ''), t.get('subject', ''), t.get('homeroom', ''), t.get('notes', '')])
    else:
        ws.append([1, "김민수", "국어", "1-1", "※ 담임학급은 '1-1' 과 같은 형태로 입력"])

    max_r = ws.max_row
    for col, width in enumerate([6, 12, 12, 12, 64], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=1, max_row=max_r, min_col=1, max_col=5):
        for cell in row:
            cell.border = thin_border
            cell.alignment = center_align
            if cell.row == 1:
                cell.font = Font(bold=True)
                cell.fill = header_fill
    wb.save(path)
    return {"success": True, "path": path}


def read_teachers(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    teachers = []
    has_number_column = False
    for row in ws.iter_rows(values_only=True):
        if not row: continue
        if str(row[0]).strip() == '순':
            has_number_column = True
            continue
        if has_number_column:
            row = row[1:]
        if not row or not row[0]: continue
        val_0 = str(row[0]).strip()
        if "안내:" in val_0 or val_0 in ("성명", "이름"): continue

        homeroom_raw = str(row[2]).strip() if len(row)>2 and row[2] and str(row[2]) != "None" else ""
        homeroom_clean = ""
        if homeroom_raw:
            m = re.search(r'(\d+)\s*[-/학년]+\s*(\d+)', homeroom_raw)
            if m: homeroom_clean = f"{m.group(1)}-{m.group(2)}"
            else: homeroom_clean = homeroom_raw

        subject = str(row[1]).strip() if len(row)>1 and row[1] and str(row[1]) != "None" else ""
        notes = str(row[3]).strip() if len(row)>3 and row[3] and str(row[3]) != "None" else ""

        teachers.append({
            "name": val_0, "subject": subject, "homeroom": homeroom_clean, "notes": notes,
            "isMain": True, "isSub": True, "isClass1": True, "isClass2": True, "isClass3": True, "isHallway": False, "isSeparated": False, "isSpecial": False
        })
    wb.close()
    return teachers


# 교사별 시트의 감독 유형 칸 순서: (role_key 값, 머리글). 모든 교사가 0회인 칸은 내보내지 않는다.
RESULT_ROLES = [('정', '정'), ('부', '부'), ('복도', '복'), ('분리', '분'), ('특수', '특'), ('수업', '수')]


def write_result(appData, path):

    wb = Workbook()

    # --- 엑셀 스타일 정의 ---
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    center_align = Alignment(horizontal='center', vertical='center')
    center_align_wrap = Alignment(horizontal='center', vertical='center', wrap_text=True)  # 줄바꿈용 서식
    header_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    periods_sorted = sorted(appData.get('periods', []), key=lambda x: (x['date'], period_num(x['name'])))
    periods_by_date = {}
    for p in periods_sorted: periods_by_date.setdefault(p['date'], []).append(p)
    dates = sorted(list(periods_by_date.keys()))
    days_kr = ['월','화','수','목','금','토','일']

    def add_header(ws, fixed, period_title):
        """앞 고정 칸(두 줄 병합) + 날짜(첫 줄 병합) / 교시(둘째 줄) 머리글."""
        row1, row2 = list(fixed), [""] * len(fixed)
        for d in dates:
            dt = datetime.strptime(d, '%Y-%m-%d')
            row1.extend([f"{dt.month}월 {dt.day}일({days_kr[dt.weekday()]})"] + [""] * (len(periods_by_date[d]) - 1))
            row2.extend(period_title(p) for p in periods_by_date[d])
        ws.append(row1); ws.append(row2)
        for i in range(1, len(fixed) + 1):
            ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
        cidx = len(fixed) + 1
        for d in dates:
            pl = len(periods_by_date[d])
            if pl > 1: ws.merge_cells(start_row=1, start_column=cidx, end_row=1, end_column=cidx + pl - 1)
            cidx += pl

    def style(ws, align):
        for r in ws.iter_rows():
            for c in r:
                c.border = thin_border
                c.alignment = align
                c.number_format = '@'
                if c.row <= 2:
                    c.font = Font(bold=True)
                    c.fill = header_fill

    teachers, assigns = appData.get('teachers', []), appData.get('assignments', {})

    # ── 교사별 시간표: 순, 성명, 정·부·복·분·특·수(0회뿐인 칸 제외), 합 ──
    ws1 = wb.active
    ws1.title = "교사별 시간표"
    counts = []
    for t in teachers:
        cnt = {k: 0 for k, _ in RESULT_ROLES}
        for p in periods_sorted:
            a = assigns.get(t['id'], {}).get(p['id'], "")
            if a and a != '-':
                cnt[role_key(a)] += 1
        counts.append(cnt)
    roles = [(k, lbl) for k, lbl in RESULT_ROLES if any(cnt[k] for cnt in counts)]
    fixed = ["순", "성명"] + [lbl for _, lbl in roles] + ["합"]
    add_header(ws1, fixed, lambda p: p['name'])
    for i, (t, cnt) in enumerate(zip(teachers, counts)):
        trow = [i + 1, t['name']] + [cnt[k] for k, _ in roles] + [sum(cnt.values())]
        trow += [assigns.get(t['id'], {}).get(p['id'], "-") for p in periods_sorted]
        ws1.append(trow)
    for col in range(1, ws1.max_column + 1):
        ws1.column_dimensions[get_column_letter(col)].width = 16 if col > len(fixed) else 5 if col == 1 else 12 if col == 2 else 4.5
    style(ws1, center_align)

    # ── 학년별 시간표: 시험 보는 학년마다 한 시트. 교시 머리글에 그 학년 과목을 괄호로 붙인다 ──
    cAssign = {}
    for t_id, p_dict in assigns.items():
        tObj = next((t for t in teachers if t['id'] == t_id), None)
        if not tObj: continue
        for p_id, a_label in p_dict.items():
            if not a_label or a_label == '-': continue
            cName, roleStr = parse_label(a_label)
            cAssign.setdefault(cName, {}).setdefault(p_id, []).append(f"{tObj['name']}({roleStr})" if roleStr else tObj['name'])

    def natural_sort_key(s):
        return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', s)]

    def get_sort_key(name_str):
        if '(정)' in name_str:
            pri = 1
        elif '(부)' in name_str:
            pri = 2
        elif '(수업)' in name_str or '(수)' in name_str:
            pri = 3
        else:
            m = re.search(r'\((\d+)\)', name_str)
            pri = 10 + int(m.group(1)) if m else 99
        return (pri, name_str)

    class_subjects = appData.get('classSubjects', {})
    for g in sorted(appData.get('config', {}).get('grades', []), key=int):
        g_str = str(g)
        ws = wb.create_sheet(f"{g}학년 시간표")

        def period_title(p):
            subj = class_subjects.get(p['id'], {}).get(g_str, '')
            return f"{p['name']}({SUBJECT_LABEL.get(subj, subj)})" if subj else p['name']

        add_header(ws, ["학급"], period_title)
        for cName in sorted((c for c in cAssign if grade_of(c) == g_str), key=natural_sort_key):
            row_data = [cName]
            for p in periods_sorted:
                assigned_list = sorted(cAssign[cName].get(p['id'], []), key=get_sort_key)
                # 한 칸에 여러 명이면 줄바꿈으로 잇는다
                row_data.append("\n".join(assigned_list) if assigned_list else "-")
            ws.append(row_data)
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20 if col > 1 else 16
        style(ws, center_align_wrap)

    wb.save(path)
