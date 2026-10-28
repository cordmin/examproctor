"""감독 명단 양식의 열 너비와 새/기존 제목 가져오기 검사."""
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from excel import read_teachers, write_template


def test_teacher_template():
    teachers = [{"name": "남궁민수", "subject": "정보", "homeroom": "1-10", "notes": "참고"},
                {"name": "김민수", "subject": "국어", "homeroom": "", "notes": ""}]
    with TemporaryDirectory() as folder:
        path = Path(folder) / "teachers.xlsx"
        write_template(path, teachers)
        wb = load_workbook(path)
        ws = wb.active
        assert [cell.value for cell in ws[1]] == ["순", "이름", "담당교과", "담임학급", "참고사항"]
        assert [ws.column_dimensions[c].width for c in "BCD"] == [12, 12, 12]
        assert all(ws[c].alignment.horizontal == "center" for c in ("A1", "B1", "C1", "A2", "B2", "C2"))
        assert [ws.cell(row, 1).value for row in range(2, 4)] == [1, 2]
        for numbered in (True, False):
            if not numbered:
                ws.delete_cols(1)
            for header in ("이름", "성명"):
                ws.cell(1, 2 if numbered else 1, header)
                wb.save(path)
                actual = read_teachers(path)
                assert len(actual) == len(teachers)
                for original, imported in zip(teachers, actual):
                    assert all(imported[key] == value for key, value in original.items())
        wb.close()
        write_template(path)
        assert len(read_teachers(path)) == 1


if __name__ == "__main__":
    test_teacher_template()
    print("PASS: template headers, equal compact widths, new/legacy import, empty-list example")
